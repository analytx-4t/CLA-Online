import docx, os

doc_path = r"c:\Users\hp\Desktop\Legal\CLA_Online_Before_After_Comparison.docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

questions = []
for p in doc.paragraphs:
    text = p.text.strip()
    if text:
        print("P:", text[:100])

for t_idx, table in enumerate(doc.tables):
    print(f"\n--- TABLE {t_idx+1} ({len(table.rows)} rows) ---")
    for r_idx, row in enumerate(table.rows):
        row_cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"Row {r_idx}: {row_cells[:3]}")
